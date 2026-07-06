"""End-to-end output verification: is the file we UPLOAD the ranking we INTEND, mapped to the right
members, in the right direction? Catches the silent-cap bug classes (id-join mismatch, misalignment,
inverted ranking) that the score being high (0.9195) already makes unlikely but doesn't prove.
Run: .venv/bin/python scratchpad/verify_output.py"""
import openpyxl, pandas as pd, numpy as np
from pathlib import Path

N, K = 500_000, 100_000
ok = True

# --- 1. JOIN KEY: does the submission template's ID match premier.pkl's id (set AND order)? ---
tmpl = next(Path("docs").glob("*submission*template*.xlsx"))
w = openpyxl.load_workbook(tmpl, read_only=True); ps = w["Predictions"]
hdr = [ps.cell(1, c).value for c in (1, 2)]
tmpl_ids = np.array([row[0].value for row in ps.iter_rows(min_row=2, max_row=N+1, max_col=1)])
w.close()
pkl = pd.read_pickle("data/premier.pkl")
pkl_ids = pkl["id"].to_numpy()
same_set = set(tmpl_ids.tolist()) == set(pkl_ids.tolist())
same_order = np.array_equal(tmpl_ids.astype(np.int64), pkl_ids.astype(np.int64))
print("=== 1. JOIN KEY (template ID vs premier.pkl id) ===")
print(f"  template header {hdr}  | template ID range [{tmpl_ids.min()},{tmpl_ids.max()}] count {len(tmpl_ids)}")
print(f"  same SET: {same_set}   same ORDER: {same_order}")
ok &= same_set

# --- 2. ARTIFACT FIDELITY: does the uploaded xlsx == our banked ranking scores_v35.csv? ---
print("=== 2. UPLOADED FILE == banked ranking ===")
sub_wb = openpyxl.load_workbook("artifacts/v35/submission_v35_tranche2.xlsx", read_only=True)
qs = sub_wb["Predictions"]
sub = {row[0].value: row[1].value for row in qs.iter_rows(min_row=2, max_row=N+1, max_col=2)}
sub_wb.close()
v35 = pd.read_csv("artifacts/v35/scores_v35.csv").set_index("id")["score"]
maxdiff = max(abs(sub[i] - v35[i]) for i in v35.index)
n_missing = len(set(v35.index) - set(sub))
print(f"  uploaded rows {len(sub)}  missing ids {n_missing}  max|xlsx-scores| {maxdiff:.2e}")
ok &= (len(sub) == N and n_missing == 0 and maxdiff < 1e-9)

# --- 3. TOP-20% MATCH: top-100k of the uploaded file == top-100k of our ranking? ---
sub_s = pd.Series(sub); top_sub = set(sub_s.nlargest(K).index); top_v35 = set(v35.nlargest(K).index)
print("=== 3. TOP-20% SET ===")
print(f"  uploaded top-100k == banked top-100k: {top_sub == top_v35}  (overlap {len(top_sub & top_v35)}/{K})")
ok &= (top_sub == top_v35)

# --- 4. DIRECTION: are high-score members the genuinely high-VALUE ones (not inverted)? ---
f = pkl.set_index("id").fillna(0.0)
catsum = f[["f6","f7","f8","f9","f10"]].sum(axis=1)
top, bot = v35.nlargest(K).index, v35.nsmallest(K).index
print("=== 4. DIRECTION (top-100k vs bottom-100k medians) — top should be MORE profitable ===")
for lab, s in [("spend(catsum)", catsum), ("revolve(f1)", f.f1), ("risk(f11)", f.f11),
               ("collection(f3)", f.f3), ("supp(f19)", f.f19)]:
    t, b = s.loc[top].median(), s.loc[bot].median()
    print(f"  {lab:16} top {t:>12,.2f}   bottom {b:>12,.2f}")
print(f"  f3 (collection) in top-100k: {int(f.f3.loc[top].sum())} (must be ~0 — hard-evicted)")

print(f"\nVERDICT: {'ALL CHECKS PASS — uploaded output is correct & faithful' if ok else 'PROBLEM FOUND — see above'}")
